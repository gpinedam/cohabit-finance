import io
from collections import defaultdict
from decimal import Decimal

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.dependencies.auth import get_current_user, get_db
from app.models.expense import Expense
from app.models.user import User
from app.services.balance_service import get_balance
from app.services.report_service import get_history, get_monthly, get_personal_summary, get_personal_tracker, _check_membership

router = APIRouter(prefix="/reports", tags=["reports"])


@router.get("/balance")
def balance(
    couple_id: int = Query(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return get_balance(db, couple_id, current_user.id)


@router.get("/monthly")
def monthly(
    couple_id: int = Query(...),
    year: int = Query(...),
    month: int = Query(..., ge=1, le=12),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return get_monthly(db, couple_id, year, month, current_user.id)


@router.get("/history")
def history(
    couple_id: int = Query(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return get_history(db, couple_id, current_user.id)


@router.get("/personal-summary")
def personal_summary(
    couple_id: int = Query(...),
    year: int = Query(...),
    month: int = Query(..., ge=1, le=12),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return get_personal_summary(db, couple_id, year, month, current_user.id)


@router.get("/personal-tracker")
def personal_tracker(
    couple_id: int = Query(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return get_personal_tracker(db, couple_id, current_user.id)


@router.get("/personal-tracker/export")
def personal_tracker_export(
    couple_id: int = Query(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _check_membership(db, couple_id, current_user.id)
    tracker = get_personal_tracker(db, couple_id, current_user.id)
    user = db.query(User).filter(User.id == current_user.id).first()
    data = _build_tracker_excel(tracker, user.name if user else "Usuario")
    headers = {"Content-Disposition": f'attachment; filename="cohabit-personal-{current_user.id}.xlsx"'}
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


# ── Excel export ─────────────────────────────────────────────────────────────

MONTH_NAMES_ES = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
]

_HEADER_FILL   = PatternFill("solid", fgColor="4F46E5")
_HEADER_FONT   = Font(bold=True, color="FFFFFF", size=10)
_SUMMARY_FILL  = PatternFill("solid", fgColor="EEF2FF")
_SUMMARY_FONT  = Font(bold=True, size=10)
_TOTAL_FONT    = Font(bold=True, size=10)
_CENTER        = Alignment(horizontal="center", vertical="center")
_RIGHT         = Alignment(horizontal="right")


def _col_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _build_tracker_excel(tracker: list[dict], user_name: str) -> bytes:
    """Build an Excel workbook for the personal financial tracker."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Historial financiero"

    headers = [
        "Mes", "Año",
        "Ingreso base (S/)", "Ingreso extra (S/)", "Total ingresos (S/)",
        "Gastos compartidos (S/)", "Gastos personales (S/)",
        "Ahorro reservado (S/)", "Fondo emergencia (S/)",
        "Disponible (S/)",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    for row_idx, m in enumerate(reversed(tracker), 2):
        disponible = float(m["available"])
        avail_cell_color = "D1FAE5" if disponible >= 0 else "FEE2E2"  # green / red tint

        ws.cell(row=row_idx, column=1, value=MONTH_NAMES_ES[m["month"] - 1])
        ws.cell(row=row_idx, column=2, value=m["year"])
        for col, key in [
            (3, "income_base"), (4, "extra_income"), (5, "income_total"),
            (6, "shared_spent"), (7, "private_spent"),
            (8, "savings_reserved"), (9, "emergency_reserved"),
        ]:
            cell = ws.cell(row=row_idx, column=col, value=float(m[key]))
            cell.number_format = '"S/ "#,##0.00'
            cell.alignment = _RIGHT
        avail_cell = ws.cell(row=row_idx, column=10, value=disponible)
        avail_cell.number_format = '"S/ "#,##0.00'
        avail_cell.alignment = _RIGHT
        avail_cell.fill = PatternFill("solid", fgColor=avail_cell_color)
        avail_cell.font = Font(bold=True, size=10)

    widths = [14, 7, 18, 18, 18, 22, 20, 20, 20, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Add user name + generation date in row below data
    info_row = len(tracker) + 3
    ws.cell(row=info_row, column=1, value=f"Usuario: {user_name}").font = Font(italic=True, size=9, color="94A3B8")
    from datetime import datetime
    ws.cell(row=info_row, column=3, value=f"Generado: {datetime.utcnow().strftime('%d/%m/%Y')}").font = Font(italic=True, size=9, color="94A3B8")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _build_excel(
    db: Session,
    couple_id: int,
    filter_year: int | None = None,
    filter_month: int | None = None,
) -> bytes:
    # ── fetch expenses ordered by date (with optional filters) ──────
    q = db.query(Expense).filter(Expense.couple_id == couple_id)
    expenses = q.order_by(Expense.created_at).all()
    if filter_year:
        expenses = [e for e in expenses if e.created_at.year == filter_year]
    if filter_month:
        expenses = [e for e in expenses if e.created_at.month == filter_month]

    # map user_id → name
    user_ids = {e.paid_by for e in expenses}
    users = {u.id: u.name for u in db.query(User).filter(User.id.in_(user_ids)).all()}

    # group by (year, month)
    months: dict[tuple, list] = defaultdict(list)
    for e in expenses:
        months[(e.created_at.year, e.created_at.month)].append(e)

    wb = Workbook()

    # ── Summary sheet ────────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Resumen"

    sum_headers = ["Mes", "Año", "Gastos", "Total (S/)"]
    for col, h in enumerate(sum_headers, 1):
        cell = ws_sum.cell(row=1, column=col, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
    ws_sum.row_dimensions[1].height = 20

    grand_total = Decimal("0")
    for row_idx, ((year, month), exps) in enumerate(sorted(months.items(), reverse=True), 2):
        month_total = sum(Decimal(str(e.total_amount)) for e in exps)
        grand_total += month_total
        ws_sum.cell(row=row_idx, column=1, value=MONTH_NAMES_ES[month - 1])
        ws_sum.cell(row=row_idx, column=2, value=year)
        ws_sum.cell(row=row_idx, column=3, value=len(exps))
        amt_cell = ws_sum.cell(row=row_idx, column=4, value=float(round(month_total, 2)))
        amt_cell.number_format = '"S/ "#,##0.00'

    # grand total row
    last = len(months) + 2
    ws_sum.cell(row=last, column=1, value="TOTAL").font = _TOTAL_FONT
    tot_cell = ws_sum.cell(row=last, column=4, value=float(round(grand_total, 2)))
    tot_cell.number_format = '"S/ "#,##0.00'
    tot_cell.font = _TOTAL_FONT
    for col in range(1, 5):
        ws_sum.cell(row=last, column=col).fill = _SUMMARY_FILL

    _col_widths(ws_sum, [16, 7, 10, 14])

    # ── One sheet per month (most recent first) ──────────────────────
    COLS = ["Fecha", "Categoría", "Subcategoría", "Descripción", "Pagado por", "Total (S/)", "Reparto"]
    WIDTHS = [14, 16, 16, 28, 14, 14, 14]

    for (year, month), exps in sorted(months.items(), reverse=True):
        sheet_name = f"{MONTH_NAMES_ES[month - 1][:3]} {year}"
        ws = wb.create_sheet(title=sheet_name)

        # header row
        for col, h in enumerate(COLS, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = _CENTER
        ws.row_dimensions[1].height = 20
        ws.freeze_panes = "A2"

        month_total = Decimal("0")
        for r, e in enumerate(sorted(exps, key=lambda x: x.created_at), 2):
            amt = Decimal(str(e.total_amount))
            month_total += amt
            ws.cell(row=r, column=1, value=e.created_at.strftime("%d/%m/%Y"))
            ws.cell(row=r, column=2, value=e.category)
            ws.cell(row=r, column=3, value=e.subcategory or "")
            ws.cell(row=r, column=4, value=e.description or "")
            ws.cell(row=r, column=5, value=users.get(e.paid_by, f"#{e.paid_by}"))
            amt_cell = ws.cell(row=r, column=6, value=float(round(amt, 2)))
            amt_cell.number_format = '"S/ "#,##0.00'
            amt_cell.alignment = _RIGHT
            ws.cell(row=r, column=7, value=e.split_type)

        # total row
        total_row = len(exps) + 2
        ws.cell(row=total_row, column=5, value="TOTAL").font = _TOTAL_FONT
        tot = ws.cell(row=total_row, column=6, value=float(round(month_total, 2)))
        tot.number_format = '"S/ "#,##0.00'
        tot.font = _TOTAL_FONT
        tot.alignment = _RIGHT
        for col in range(1, len(COLS) + 1):
            ws.cell(row=total_row, column=col).fill = _SUMMARY_FILL

        _col_widths(ws, WIDTHS)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


@router.get("/export")
def export_excel(
    couple_id: int = Query(...),
    year: int | None = Query(None),
    month: int | None = Query(None, ge=1, le=12),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _check_membership(db, couple_id, current_user.id)
    data = _build_excel(db, couple_id, year, month)
    fname = "cohabit-historial"
    if year:
        fname += f"-{year}"
        if month:
            fname += f"-{MONTH_NAMES_ES[month - 1]}"
    headers = {"Content-Disposition": f'attachment; filename="{fname}.xlsx"'}
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )
